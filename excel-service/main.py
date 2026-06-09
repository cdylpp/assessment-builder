from io import BytesIO
import json
from fastapi import FastAPI, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

# TODO: define styled templates using .xltx

app = FastAPI()

@app.get("/health")
def health():
    return {"status": "ok"}

def excel_cell_value(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value)
    return value

@app.post("/generate")
def generate(payload: dict):
    wb = Workbook()
    ws = wb.active
    ws.title = "Assessment Template"

    ws.append(["Evolution", "Metric", "Type", "Domain", "Score", "Notes"])

    for evolution in payload["event"]["evolutions"]:
        for metric in evolution["metrics"]:
            ws.append([
                excel_cell_value(evolution["name"]),
                excel_cell_value(metric["name"]),
                excel_cell_value(metric["type"]),
                excel_cell_value(metric["domain"]),
                "",
                ""
            ])
    
    last_row = ws.max_row
    table = Table(displayName="AssessmentMetrics", ref=f'A1:F{last_row}')
    style = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    table.tableStyleInfo = style
    ws.add_table(table)
    stream = BytesIO()
    wb.save(stream)
    return Response(
        stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
